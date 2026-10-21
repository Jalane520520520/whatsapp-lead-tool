#!/usr/bin/env python3
"""Patch-by-patch pipeline: pick the next N unfinished 1 km cells -> Google Maps (gosom) per cell -> master.csv
-> website WhatsApp (new businesses only) -> WhatsApp check (new phones only) -> Excel from the WHOLE master.
State lives in tool/data/<slug>/ (progress.json, master.csv, web_cache.csv, wa_cache.csv); a cell is saved the moment it finishes,
so killing a run loses nothing that was finished. Never sends messages.
"""
import argparse, csv, datetime, json, os, signal, subprocess, sys, time
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import cells as C

import paths
HERE = C.HERE
TOOL = C.TOOL
PY = sys.executable
HEAD = ["公司名称", "WhatsApp号码", "来源", "电话", "电话类型", "官网", "WhatsApp出处链接", "谷歌地图链接", "地址", "类目"]
MASTER_COLS = ["name", "phone_e164", "phone_type", "website", "address", "category", "latitude", "longitude",
               "maps_link", "place_id", "city", "cell_id"]
WEB_COLS = ["place_id", "wa_number_e164", "wa_source_url", "wa_evidence", "site_status"]
CHILD = None   # current child process (own process group), killed on SIGTERM/SIGINT


def log(*a):
    print(time.strftime("%H:%M:%S"), *a, flush=True)


def _die(signum, frame):
    if CHILD and CHILD.poll() is None:
        try: os.killpg(CHILD.pid, signal.SIGKILL)
        except Exception: pass
    log("interrupted; finished cells are already saved")
    os._exit(130)


def run_tee(cmd, cwd=None, quiet=False):
    """Run in its own process group, echo output, return (returncode, browser-crash lines)."""
    global CHILD
    CHILD = p = subprocess.Popen(cmd, cwd=cwd, env=paths.child_env(), stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True,
                                 errors="replace", start_new_session=True)
    crashes = 0
    for line in p.stdout:
        if "SIGTRAP" in line or "Target page, context or browser has been closed" in line or "target closed" in line.lower():
            crashes += 1
        if not quiet or os.environ.get("WAT_VERBOSE"): print(line.rstrip()[:300], flush=True)
    rc = p.wait(); CHILD = None
    return rc, crashes


def read(path):
    if not os.path.exists(path): return []
    with open(path, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def append_rows(path, cols, rows):
    new = not os.path.exists(path)
    with open(path, "a", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        if new: w.writeheader()
        w.writerows(rows); f.flush(); os.fsync(f.fileno())


def first_num(s):
    import re
    m = re.findall(r"\+\d{6,15}", s or "")
    return m[0] if m else ""


# ---------------------------------------------------------------- step 1: maps, cell by cell
def crawl_cell(cell, kws, args, d, master_ids):
    """returns (status, rows_new_list, n_rows_total, warning). Retries a cell up to 3 times if Chromium crashed with 0 rows."""
    tmp = os.path.join(d, "tmp"); os.makedirs(tmp, exist_ok=True)
    got, warn = {}, ""
    for kw_i, kw in enumerate(kws):
        out = os.path.join(tmp, f"cell_{kw_i}.csv")
        best, crash_left = None, 0
        for attempt in range(1, 4):
            if os.path.exists(out): os.remove(out)
            cmd = [PY, os.path.join(TOOL, "maps", "run_maps.py"), "--keyword", kw, "--country", args.country,
                   "--bbox", ",".join(map(str, cell["bbox"])), "--cell-km", str(C.CELL_KM), "--out", out]
            rc, crashes = run_tee(cmd, cwd=os.path.join(TOOL, "maps"), quiet=True)
            rows = read(out) if rc == 0 else None
            if rows is not None and (best is None or len(rows) > len(best)):
                best, crash_left = rows, crashes
            if rc == 0 and (rows or crashes == 0):
                break
            log(f"  retry '{kw}' (attempt {attempt}/3): exit={rc}, rows={len(rows) if rows is not None else '-'}, crashes={crashes}")
        if best is None or (not best and crash_left):
            return "failed", [], 0, f"Chrome crashed on '{kw}'"
        for r in best:
            got.setdefault(r["place_id"] or r["maps_link"], r)
        if crash_left: warn = f"Chrome crashed {crash_left}x on '{kw}' but rows were found"
    new = []
    for k, r in got.items():
        if k in master_ids: continue
        master_ids.add(k)
        new.append(dict(r, place_id=r["place_id"] or k, city=cell["city"], cell_id=cell["id"]))
    return "done", new, len(got), warn


# ---------------------------------------------------------------- step 2: websites (new businesses only)
def website_step(d, country):
    master = read(os.path.join(d, "master.csv"))
    cache_p = os.path.join(d, "web_cache.csv")
    have = {r["place_id"] for r in read(cache_p)}
    todo = [r for r in master if r["place_id"] not in have]
    if not todo:
        log("[website] nothing new"); return
    log(f"[website] {len(todo)} new businesses (slow, ~10 s per site)")
    web = os.path.join(TOOL, "web")
    code = ("import sys,asyncio;sys.path.insert(0,sys.argv[1]);import find_wa;find_wa.PAGE_TIMEOUT_MS=int(sys.argv[5]);"
            "asyncio.run(find_wa.main(sys.argv[2],sys.argv[3],sys.argv[4]))")
    tmp = os.path.join(d, "tmp"); os.makedirs(tmp, exist_ok=True)
    rin, p1 = os.path.join(tmp, "web_in.csv"), os.path.join(tmp, "web_pass1.csv")
    with open(rin, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=MASTER_COLS, extrasaction="ignore"); w.writeheader(); w.writerows(todo)
    if os.path.exists(p1): os.remove(p1)
    rc, _ = run_tee([PY, "-c", code, web, rin, p1, country, "30000"], cwd=web)
    if rc != 0 and not read(p1):
        log("ERROR website step failed"); return
    rows1 = read(p1)
    bad = [r for r in rows1 if (r.get("site_status") or "").startswith("error")]
    final = {r["place_id"]: r for r in rows1}
    if bad:
        log(f"[website] retrying {len(bad)} failed sites (60 s)")
        rr, ro = os.path.join(tmp, "web_retry_in.csv"), os.path.join(tmp, "web_retry_out.csv")
        with open(rr, "w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=MASTER_COLS, extrasaction="ignore"); w.writeheader(); w.writerows(bad)
        if os.path.exists(ro): os.remove(ro)
        run_tee([PY, "-c", code, web, rr, ro, country, "60000"], cwd=web)
        for r in read(ro): final[r["place_id"]] = r
    append_rows(cache_p, WEB_COLS, list(final.values()))


# ---------------------------------------------------------------- step 3: WhatsApp check (new phones only)
def wa_step(d, skip):
    web = {r["place_id"]: r for r in read(os.path.join(d, "web_cache.csv"))}
    cache_p = os.path.join(d, "wa_cache.csv")
    have = {r["phone_e164"] for r in read(cache_p)}
    nums = sorted({r["phone_e164"] for r in read(os.path.join(d, "master.csv"))
                   if r.get("phone_e164") and r["phone_e164"] not in have
                   and not first_num(web.get(r["place_id"], {}).get("wa_number_e164"))})
    if skip:
        log(f"[whatsapp] check skipped ({len(nums)} new numbers would be checked)"); return 0
    if not nums:
        log("[whatsapp] nothing new"); return 0
    tmp = os.path.join(d, "tmp"); os.makedirs(tmp, exist_ok=True)
    nin, nout = os.path.join(tmp, "wa_in.csv"), os.path.join(tmp, "wa_out.csv")
    with open(nin, "w", encoding="utf-8", newline="") as f: f.write("phone_e164\n" + "\n".join(nums) + "\n")
    log(f"[whatsapp] checking {len(nums)} new numbers (read only, no messages sent)")
    rc, _ = run_tee([paths.NODE, "check.js", "--in", nin, "--out", nout, "--require-login"], cwd=os.path.join(TOOL, "wa_check"))
    if rc == 3:  # not logged in: the run still finishes; these phones stay 未检测
        log(f"[whatsapp] 未登录 WhatsApp，跳过验证（{len(nums)} 个电话标为未检测；登录后再跑会补查）"); return 0
    got = [{"phone_e164": r["phone_e164"], "on_whatsapp": "yes" if r["on_whatsapp"].strip().lower() in ("true", "yes", "1") else "no"}
           for r in read(nout) if r.get("on_whatsapp", "").strip()]
    append_rows(cache_p, ["phone_e164", "on_whatsapp"], got)
    if rc != 0:
        log("STOPPED: WhatsApp check did not finish (login needed or account restricted). Run again after fixing.")
    return rc


# ---------------------------------------------------------------- step 4: Excel from the whole master
PTYPE = {"mobile": "手机", "fixed_line": "座机", "fixed_line_or_mobile": "座机或手机", "voip": "网络电话", "toll_free": "免费电话"}


def table_rows(d):
    """One dict per business in the whole master, WhatsApp ones first. Used by the Excel and the window."""
    master = read(os.path.join(d, "master.csv"))
    web = {r["place_id"]: r for r in read(os.path.join(d, "web_cache.csv"))}
    wa = {r["phone_e164"]: r["on_whatsapp"] for r in read(os.path.join(d, "wa_cache.csv"))}
    rows = []
    for r in master:
        w = web.get(r["place_id"], {})
        site_wa, phone = first_num(w.get("wa_number_e164")), r.get("phone_e164", "")
        num, src = "", ""
        if site_wa: num, src = site_wa, "官网"
        elif phone and wa.get(phone) == "yes": num, src = phone, "WhatsApp验证"
        elif not phone: src = "无电话"
        else: src = "未开通" if phone in wa else "未检测"
        rows.append({"name": r["name"], "wa": num, "src": src, "phone": phone,
                     "ptype": PTYPE.get(r.get("phone_type", ""), "未知" if phone else ""),
                     "website": r.get("website", ""), "wa_url": w.get("wa_source_url", ""), "maps": r.get("maps_link", ""),
                     "address": r.get("address", ""), "category": r.get("category", ""), "city": r.get("city", ""),
                     "site_error": (w.get("site_status") or "").startswith("error")})
    rows.sort(key=lambda x: 0 if x["wa"] else 1)
    return rows


def _sheet(wb, title, head, widths, data, links):
    from openpyxl.styles import Font, PatternFill
    ws = wb.create_sheet(title); ws.append(head)
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F6F43")
    for r in data:
        ws.append([v if not isinstance(v, tuple) else "打开" if v[0] else "" for v in r])
        for ci, v in enumerate(r, 1):
            if isinstance(v, tuple) and v[0]:
                cell = ws.cell(ws.max_row, ci); cell.hyperlink = v[0]; cell.font = Font(color="0563C1", underline="single")
    for i, wd in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = wd
    ws.freeze_panes = "A2"
    if data: ws.auto_filter.ref = ws.dimensions
    return ws


def build_excel(args, d, prog, warnings):
    from openpyxl import Workbook
    rows = table_rows(d)
    wb = Workbook(); wb.remove(wb.active)
    waw = [r for r in rows if r["wa"]]
    _sheet(wb, "WhatsApp客户", ["公司名称", "WhatsApp号码", "纯数字号码", "聊天", "来源", "城市", "类目", "谷歌地图"],
           (34, 18, 16, 8, 14, 14, 24, 10),
           [[r["name"], r["wa"], r["wa"].lstrip("+"), ("https://wa.me/" + r["wa"].lstrip("+"),), r["src"], r["city"], r["category"],
             (r["maps"],)] for r in waw], True)
    _sheet(wb, "全部商家", ["公司名称", "WhatsApp号码", "来源", "电话", "电话类型", "城市", "类目", "地址", "官网", "WhatsApp出处", "谷歌地图"],
           (34, 18, 14, 18, 10, 14, 24, 44, 8, 12, 10),
           [[r["name"], r["wa"], r["src"], r["phone"], r["ptype"], r["city"], r["category"], r["address"], (r["website"],),
             (r["wa_url"],), (r["maps"],)] for r in rows], True)
    sm = C.summary(prog)
    n_web = sum(r["src"] == "官网" for r in rows); n_chk = sum(r["src"] == "WhatsApp验证" for r in rows)
    stats = [("更新时间", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")), ("关键词", "; ".join(prog["keywords"])), ("国家", args.country),
             ("已完成格数/总格数", f"{sm['done']}/{sm['total']}"),
             ("已完成城市", "、".join(sm["cities_done"]) or "（暂无）"),
             ("下次从哪个城市第几格开始", f"{sm['next_city']} 第{sm['next_n']}格" if sm["next_city"] else "全部完成"),
             ("商家数", len(rows)), ("有电话", sum(1 for r in rows if r["phone"])), ("有官网", sum(1 for r in rows if r["website"])),
             ("WhatsApp(来自官网)", n_web), ("WhatsApp(来自验证)", n_chk), ("WhatsApp总数", n_web + n_chk),
             ("官网打开失败", sum(r["site_error"] for r in rows))]
    failed = [c["id"] for c in prog["cells"] if c.get("status") == "failed"]
    if failed: stats.append(("失败的格子(下次会重试)", "、".join(failed[:20])))
    for w in warnings: stats.append(("警告", w))
    if any(r["src"] == "未检测" for r in rows):
        stats.append(("备注", "有电话还没做 WhatsApp 验证（未检测）；登录 WhatsApp 后再抓一次会补查"))
    s = wb.create_sheet("说明")
    for k, v in stats: s.append([k, v])
    s.column_dimensions["A"].width = 26; s.column_dimensions["B"].width = 60
    try:
        wb.save(args.out)
    except PermissionError:  # Windows: the client has this Excel open; keep the data, write a side copy
        side = args.out[:-5] + " (新).xlsx"
        wb.save(side); log(f"EXCEL_LOCKED: {args.out} 正在被打开，这次另存为 {side}"); args.out = side
    return {"businesses": len(rows), "excel_rows": len(rows), "whatsapp": n_web + n_chk, **{k: sm[k] for k in ("done", "total")}}


def main():
    signal.signal(signal.SIGTERM, _die); signal.signal(signal.SIGINT, _die)
    a = argparse.ArgumentParser()
    a.add_argument("--keywords", required=True, help="keyword1;keyword2")
    a.add_argument("--country", required=True)
    a.add_argument("--cells", type=int, default=3, help="how many unfinished 1 km cells to crawl this run")
    a.add_argument("--city", help="only take cells of this city")
    a.add_argument("--out", help="Excel path (default 表格/<国家> - <关键词>.xlsx)")
    a.add_argument("--run-dir", help="folder for this run's log/Excel (server uses it)")
    a.add_argument("--skip-wa-check", action="store_true")
    args = a.parse_args()
    args.country = args.country.upper()
    kws = [k.strip() for k in args.keywords.split(";") if k.strip()]
    d = C.data_dir(args.country, kws); os.makedirs(d, exist_ok=True)
    if args.run_dir: os.makedirs(args.run_dir, exist_ok=True)
    if not args.out: args.out = C.excel_path(args.country, kws); os.makedirs(C.TABLES, exist_ok=True)
    elif not os.path.isabs(args.out): args.out = os.path.join(args.run_dir or d, args.out)
    prog = C.load_progress(args.country, kws)
    if args.city and not any(c["city"].lower() == args.city.lower() for c in prog["cells"]):
        sys.exit(f"city '{args.city}' not in the {args.country} list")
    log(f"data folder: {d}")
    sm = C.summary(prog); log(f"progress before: {sm['done']}/{sm['total']} cells")
    master_p = os.path.join(d, "master.csv")
    master_ids = {r["place_id"] for r in read(master_p)}
    todo = [c for c in prog["cells"] if c.get("status") != "done" and (not args.city or c["city"].lower() == args.city.lower())][:max(0, args.cells)]
    log(f"this run: {len(todo)} cells" + (f" ({todo[0]['id']} .. {todo[-1]['id']})" if todo else ""))
    warnings = []
    for i, cell in enumerate(todo, 1):
        t0 = time.time()
        log(f"[maps {i}/{len(todo)}] {cell['id']} bbox={cell['bbox']}")
        status, new, n_total, warn = crawl_cell(cell, kws, args, d, master_ids)
        if status == "done":
            if new: append_rows(master_p, MASTER_COLS, new)       # master first, then the ledger
        else:
            for r in new: master_ids.discard(r["place_id"])
        cell.update(status=status, rows=n_total, time=time.strftime("%Y-%m-%d %H:%M:%S"), secs=round(time.time() - t0))
        C.save_progress(args.country, kws, prog)
        if warn: warnings.append(f"{cell['id']}: {warn}")
        log(f"  -> {status}: {n_total} places, {len(new)} new, {time.time() - t0:.0f}s (master now {len(master_ids)})")
    website_step(d, args.country)
    rc = wa_step(d, args.skip_wa_check)
    stats = build_excel(args, d, prog, warnings)
    log("DONE", json.dumps(stats, ensure_ascii=False)); log("EXCEL:", args.out)
    sys.exit(2 if rc else 0)


if __name__ == "__main__":
    main()
